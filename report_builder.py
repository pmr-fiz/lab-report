"""
report_builder.py — шаблон отчёта, закодированный по образцу.

Инкапсулирует стандарт оформления (A4; поля 3/1.5/2/2 см; Times New Roman 14 pt;
выравнивание по ширине; красная строка 1.25 см; интервал 1.5) и структуру:
титульник -> цель -> задания -> контрольные вопросы.

Зависит от code_block.py (тёмные блоки кода и вывода) — держи оба файла рядом.

Публичные функции:
    new_report()                         -> Document с нужными полями/стилем
    title_page(doc, ...)                 -> титульный лист МТУСИ с плейсхолдерами
    goal(doc, items)                     -> "Цель работы:"
    task_heading(doc, text)              -> "Задание N. ..."
    body(doc, text)                      -> абзац основного текста (по ширине, кр. строка)
    bullets(doc, items)                  -> маркированный список
    numbered(doc, items)                 -> нумерованный список
    answer_items(doc, items)             -> для заданий-вопросов: Ответ + Обоснование
    solution_heading(doc)                -> "Решение:"
    add_code_block / add_output_block    -> из code_block.py (реэкспорт)
    explanation(doc, paragraphs)         -> абзацы объяснения кода
    control_questions(doc, qa)           -> "Контрольные вопросы:" (вопрос жирным + ответ)
    save(doc, path)
"""

from docx import Document
from docx.shared import Pt, Cm, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION  # noqa: F401
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# тёмные блоки кода/вывода из соседнего модуля
from code_block import add_code_block, add_output_block  # noqa: F401

BODY_FONT = "Times New Roman"
BODY_SIZE = 14          # pt — основной текст
TABLE_SIZE = 12         # pt — таблицы
FIRST_LINE = 1.25       # см — красная строка
LINE_SPACING = 1.5


def _force_font(run, name=BODY_FONT):
    """Закрепляем шрифт для всех наборов (в т.ч. кириллицы)."""
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), name)


def new_report():
    """Пустой документ с полями и базовым стилем по стандарту МТУСИ."""
    doc = Document()

    # --- страница A4 + поля ---
    sec = doc.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.left_margin = Cm(3.0)
    sec.right_margin = Cm(1.5)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)

    # --- стиль Normal: TNR 14, интервал 1.5 ---
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(BODY_SIZE)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), BODY_FONT)
    pf = normal.paragraph_format
    pf.line_spacing = LINE_SPACING
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    return doc


# ---------- титульный лист ----------

def _center(doc, text, bold=False, size=BODY_SIZE, after=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    _force_font(r)
    return p


def _blank(doc, n=1):
    for _ in range(n):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)


def title_page(doc,
               lab_number="N",
               lab_title="НАЗВАНИЕ РАБОТЫ",
               department="{{DEPARTMENT}}",
               group="{{GROUP}}",
               student_name="{{STUDENT_NAME}}",
               city="Москва",
               year="2026"):
    """Титульный лист. Незаполненные поля остаются плейсхолдерами {{...}}."""
    _center(doc, "МИНИСТЕРСТВО ЦИФРОВОГО РАЗВИТИЯ, СВЯЗИ И МАССОВЫХ "
                 "КОММУНИКАЦИЙ РОССИЙСКОЙ ФЕДЕРАЦИИ", bold=True)
    _center(doc, "Ордена трудового Красного Знамени федеральное "
                 "государственное бюджетное образовательное учреждение "
                 "высшего образования", bold=True)
    _center(doc, "«Московский технический университет связи и информатики»",
            bold=True, after=6)
    _blank(doc, 1)
    _center(doc, f"Кафедра {department}")
    _blank(doc, 5)
    _center(doc, f"Отчет по лабораторной работе №{lab_number}.", bold=True)
    _center(doc, f"«{lab_title}».", bold=True)
    _blank(doc, 3)

    # выполнил — по правому краю, как принято на титульниках
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run(f"Выполнил: студент группы {group}")
    _force_font(r)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p2.paragraph_format.first_line_indent = Cm(0)
    r2 = p2.add_run(student_name)
    _force_font(r2)

    _blank(doc, 6)
    _center(doc, f"{city}, {year}")

    # чистый переход на новую страницу (без пустой страницы-призрака)
    new_sec = doc.add_section(WD_SECTION.NEW_PAGE)
    new_sec.page_width = Mm(210)
    new_sec.page_height = Mm(297)
    new_sec.left_margin = Cm(3.0)
    new_sec.right_margin = Cm(1.5)
    new_sec.top_margin = Cm(2.0)
    new_sec.bottom_margin = Cm(2.0)


# ---------- разделы тела ----------

def _heading(doc, text, before=10, after=4):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.bold = True
    _force_font(r)
    return p


def body(doc, text, justify=True, indent=True):
    """Абзац основного текста: по ширине, с красной строкой."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if justify else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(FIRST_LINE if indent else 0)
    r = p.add_run(text)
    _force_font(r)
    return p


def goal(doc, items):
    """'Цель работы:' + пункты (строка или список строк)."""
    _heading(doc, "Цель работы:")
    if isinstance(items, str):
        items = [items]
    for i, it in enumerate(items, 1):
        body(doc, f"{i}. {it}" if len(items) > 1 else it)


def task_heading(doc, text):
    """'Задание N. Название'."""
    _heading(doc, text)


def bullets(doc, items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(it)
        _force_font(r)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def numbered(doc, items):
    for it in items:
        p = doc.add_paragraph(style="List Number")
        r = p.add_run(it)
        _force_font(r)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def answer_items(doc, items):
    """Для заданий-вопросов. items = [{'q':..., 'answer':..., 'rationale':...}]."""
    for i, it in enumerate(items, 1):
        if it.get("q"):
            body(doc, f"{i}. {it['q']}")
        # Ответ
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(FIRST_LINE)
        rb = p.add_run("Ответ: ")
        rb.bold = True
        _force_font(rb)
        r = p.add_run(it["answer"])
        _force_font(r)
        # Обоснование
        if it.get("rationale"):
            p2 = doc.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p2.paragraph_format.first_line_indent = Cm(FIRST_LINE)
            rb2 = p2.add_run("Обоснование: ")
            rb2.bold = True
            _force_font(rb2)
            r2 = p2.add_run(it["rationale"])
            _force_font(r2)


def solution_heading(doc):
    _heading(doc, "Решение:", before=8, after=4)


def explanation(doc, paragraphs):
    """Абзацы объяснения кода (строка или список строк)."""
    if isinstance(paragraphs, str):
        paragraphs = [paragraphs]
    for para in paragraphs:
        body(doc, para)


def control_questions(doc, qa):
    """'Контрольные вопросы:' + [{'q':..., 'a':...}]."""
    _heading(doc, "Контрольные вопросы:")
    for i, item in enumerate(qa, 1):
        pq = doc.add_paragraph()
        pq.paragraph_format.first_line_indent = Cm(0)
        pq.paragraph_format.space_before = Pt(6)
        rq = pq.add_run(f"{i}. {item['q']}")
        rq.bold = True
        _force_font(rq)
        body(doc, item["a"])


def save(doc, path):
    doc.save(path)
